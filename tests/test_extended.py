import io
import os
import tarfile
import tempfile
import unittest
import zipfile
from pathlib import Path
from PIL import Image
from filetypes import inspect_file, common_actions, backend
from converter import convert
from extended import soffice

class ExtendedTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def test_extended_images(self):
        source=self.root/'photo.png'; Image.new('RGBA',(32,32),(50,100,200,80)).save(source)
        actions=inspect_file(source).actions
        self.assertTrue({'PNG','JPG','WEBP','HEIC','TIFF','AVIF','BMP','PDF'} <= set(actions))
        for action in actions:
            with self.subTest(action=action):
                out=convert(source,action)
                self.assertGreater(out.stat().st_size,0)
                self.assertEqual(inspect_file(out).kind, 'PDF' if action=='PDF' else 'SVG' if action=='SVG' else 'Bild')
        svg=self.root/'vector.svg'; svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="32" height="32"><circle cx="16" cy="16" r="12" fill="red"/></svg>')
        self.assertEqual(inspect_file(svg).kind,'SVG')
        with Image.open(convert(svg,'PNG')) as image: self.assertEqual(image.size,(32,32))
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="https://example.org/x.png"/></svg>')
        self.assertEqual(inspect_file(svg).actions,())

    def test_subtitle_roundtrip(self):
        source=self.root/'test.srt'; source.write_text('1\n00:00:01,200 --> 00:00:02,300\nHallo Welt\n')
        self.assertEqual(inspect_file(source).kind,'SRT')
        vtt=convert(source,'VTT'); self.assertTrue(vtt.read_text().startswith('WEBVTT'))
        self.assertIn('00:00:01,200',convert(vtt,'SRT').read_text())
        self.assertEqual(convert(source,'TXT').read_text(),'Hallo Welt')
        vtt.write_text('WEBVTT\n\n00:01.200 --> 00:02.300 align:start\nKurz\n')
        self.assertEqual(inspect_file(vtt).kind,'VTT')
        self.assertIn('00:00:01,200',convert(vtt,'SRT').read_text())

    def test_archives_and_traversal(self):
        source=self.root/'data.zip'
        with zipfile.ZipFile(source,'w') as z: z.writestr('folder/hello.txt','hello')
        for action in ['ZIP','TAR','GZIP']:
            out=convert(source,action)
            if action=='ZIP':
                with zipfile.ZipFile(out) as z: self.assertEqual(z.read('folder/hello.txt'),b'hello')
            else:
                with tarfile.open(out) as t: self.assertEqual(t.extractfile('folder/hello.txt').read(),b'hello')
                zipout=convert(out,'ZIP')
                with zipfile.ZipFile(zipout) as z: self.assertEqual(z.read('folder/hello.txt'),b'hello')
        evil=self.root/'evil.zip'
        with zipfile.ZipFile(evil,'w') as z:z.writestr('../escape.txt','bad')
        with self.assertRaisesRegex(ValueError,'Unsicherer'):
            convert(evil,'TAR')
        self.assertFalse((self.root/'escape.txt').exists())
        self.assertFalse((self.root/'evil_converted.tar').exists())

    def test_pdf_exports_all_pages_and_tables(self):
        import pymupdf
        source=self.root/'sample.pdf'
        doc=pymupdf.open()
        for index in range(2):
            page=doc.new_page(width=200,height=200)
            page.insert_text((20,30),'Page '+str(index+1))
            for x in [20,100,180]:page.draw_line((x,60),(x,120))
            for y in [60,90,120]:page.draw_line((20,y),(180,y))
            page.insert_text((25,80),'Name');page.insert_text((105,80),'Value')
            page.insert_text((25,110),'Test');page.insert_text((105,110),'42')
        doc.save(source);doc.close()
        for action in inspect_file(source).actions:
            with self.subTest(action=action):
                out=convert(source,action)
                self.assertGreater(out.stat().st_size,0)
                if action in ('PNG','JPG'):
                    with zipfile.ZipFile(out) as z:
                        self.assertEqual(len(z.namelist()),2)
                        with Image.open(io.BytesIO(z.read(z.namelist()[0]))) as image:self.assertEqual(image.size,(834,834))
                elif action=='PPTX':
                    from pptx import Presentation
                    self.assertEqual(len(Presentation(out).slides),2)
                elif action=='DOCX':
                    from docx import Document
                    self.assertIn('Page 1','\n'.join(p.text for p in Document(out).paragraphs))
                elif action=='XLSX':
                    from openpyxl import load_workbook
                    book=load_workbook(out);self.assertEqual(len(book.sheetnames),2);book.close()
        scanned=self.root/'scan.pdf'
        doc=pymupdf.open();doc.new_page();doc.save(scanned);doc.close()
        with self.assertRaisesRegex(ValueError,'OCR'):convert(scanned,'TXT')

    @unittest.skipUnless(backend('pandoc'),'Pandoc missing')
    def test_document_and_ebook_conversions(self):
        source=self.root/'note.md';source.write_text('# Test\n\nHello from OrbitDrop\n')
        for action in ['DOCX','MD','TXT','EPUB','HTML']:
            with self.subTest(action=action):
                out=convert(source,action);self.assertGreater(out.stat().st_size,0)
                if action in ('EPUB','DOCX'):
                    self.assertEqual(inspect_file(out).kind,action)
                    self.assertIn('Hello',convert(out,'TXT').read_text())

    @unittest.skipUnless(soffice(),'LibreOffice missing')
    def test_office_sheet_and_slides(self):
        from openpyxl import Workbook,load_workbook
        source=self.root/'sheet.xlsx'
        book=Workbook();book.active.append(['Name','Value']);book.active.append(['Test',42]);book.save(source)
        for action in ['ODS','CSV','TSV','PDF','XLSX']:
            with self.subTest(action=action):
                out=convert(source,action);self.assertGreater(out.stat().st_size,0)
                if action in ('CSV','TSV'):self.assertIn('Test',out.read_text(encoding='utf-8'))
        from pptx import Presentation
        slidefile=self.root/'slides.pptx';deck=Presentation();deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text='OrbitDrop';deck.save(slidefile)
        for action in ['PDF','PNG','JPG']:
            out=convert(slidefile,action);self.assertGreater(out.stat().st_size,0)

if __name__=='__main__':unittest.main()
