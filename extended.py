"""Additional local format backends. No automatic downloads or network conversions."""
import csv
import gzip
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import sys
import tarfile
import tempfile
import zipfile
from PIL import Image, ImageOps
from task_runtime import checkpoint, copy_stream
from filetypes import FileInfo, backend, run

IMAGE_TYPES = {'PNG': '.png', 'JPG': '.jpg', 'WEBP': '.webp', 'HEIC': '.heic',
               'TIFF': '.tiff', 'AVIF': '.avif', 'BMP': '.bmp', 'PDF': '.pdf'}

def image_actions():
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
    except ImportError:
        pass
    Image.init()
    result = []
    for target in IMAGE_TYPES:
        name = {'JPG':'JPEG', 'HEIC':'HEIF'}.get(target, target)
        if name in Image.SAVE:
            result.append(target)
    return tuple(result)

def soffice():
    found = backend('soffice')
    if found:
        console=Path(found).with_suffix('.com')
        return str(console) if sys.platform=='win32' and console.is_file() else found
    for env in ['PROGRAMFILES', 'PROGRAMFILES(X86)']:
        folder = Path(os.environ.get(env, 'C:/Program Files')) / 'LibreOffice/program'
        for name in (('soffice.com','soffice.exe') if sys.platform=='win32' else ('soffice.exe',)):
            candidate=folder/name
            if candidate.is_file():return str(candidate)
    return None


def document_actions(kind):
    pandoc, office = backend('pandoc'), soffice()
    if kind in ('XLSX', 'ODS', 'XLS', 'CSV', 'TSV'):
        if not office:
            return ()
        return ('XLSX', 'ODS', 'CSV', 'TSV', 'PDF')
    if kind in ('PPT', 'PPTX'):
        return ('PDF', 'JPG', 'PNG') if office else ()
    actions = []
    if pandoc and kind in ('DOCX', 'EPUB', 'MD', 'TXT', 'HTML'):
        actions += ['DOCX', 'MD', 'TXT', 'EPUB', 'HTML']
    elif office and kind in ('DOCX', 'RTF', 'DOC', 'PAGES', 'TXT', 'HTML'):
        actions += ['DOCX', 'TXT', 'HTML', 'EPUB']
    if office and kind != 'EPUB':
        actions.append('PDF')
    return tuple(dict.fromkeys(actions))

def svg_bytes(path):
    from defusedxml import ElementTree
    data = Path(path).read_bytes()
    root = ElementTree.fromstring(data)
    if root.tag.rsplit('}',1)[-1] != 'svg':
        return False
    # Keep SVG conversion entirely local and self-contained.
    for node in root.iter():
        if node.tag.rsplit('}',1)[-1] in ('image', 'script', 'foreignObject'):
            raise ValueError('SVG mit eingebetteten Bildern/Skripten noch nicht unterstützt.')
        for key, value in node.attrib.items():
            if key.rsplit('}',1)[-1] == 'href' and not value.startswith('#'):
                raise ValueError('Externe SVG-Verweise werden nicht geladen.')
    return data

def identify_extra(path, header):
    path = Path(path)
    if header.lstrip().startswith((b'<', b'<?xml')):
        try:
            data = svg_bytes(path)
            if data:
                return FileInfo('SVG', image_actions(), 'SVG wird gerastert; keine SVG-Ausgabe.')
        except Exception as error:
            if b'<svg' in header:
                return FileInfo('SVG', (), str(error))
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if 'word/document.xml' in names:
                kind = 'DOCX'
            elif 'ppt/presentation.xml' in names:
                kind = 'PPTX'
            elif 'xl/workbook.xml' in names:
                kind = 'XLSX'
            elif 'mimetype' in names:
                with archive.open('mimetype') as handle:
                    mime = handle.read(150)
                kind = 'EPUB' if mime == b'application/epub+zip' else 'ODS' if mime == b'application/vnd.oasis.opendocument.spreadsheet' else 'ZIP'
            elif any(n.startswith('Index/') and n.endswith('.iwa') for n in names) and path.suffix.lower() == '.pages':
                kind = 'PAGES'
            else:
                kind = 'ZIP'
            if kind != 'ZIP':
                return FileInfo(kind, document_actions(kind), 'Office: LibreOffice; Text/E-Books: Pandoc. Layout kann abweichen.')
            if any(i.flag_bits & 1 for i in archive.infolist()):
                return FileInfo('ZIP', (), 'Verschlüsseltes Archiv noch nicht unterstützt.')
            return FileInfo('ZIP', ('ZIP', 'TAR', 'GZIP'), 'Archivinhalt wird neu verpackt, nicht nur umbenannt.')
    if header.startswith(b'Rar!\x1a\x07'):
        try:
            import rarfile
            with rarfile.RarFile(path) as archive:
                encrypted = archive.needs_password()
            enabled = bool(backend('unrar') or backend('unar') or backend('bsdtar'))
            return FileInfo('RAR', ('ZIP', 'TAR', 'GZIP') if enabled and not encrypted else (), 'RAR nur Eingabe; benötigt unrar/unar/bsdtar; keine verschlüsselten Archive.')
        except ImportError:
            return FileInfo('RAR', (), 'RAR-Modul fehlt.')
    if header.startswith(b'\x1f\x8b'):
        return FileInfo('GZIP', ('ZIP', 'TAR', 'GZIP'), 'GZIP-Ausgabe ist ein TAR.GZ-Archiv.')
    if len(header) > 262 and header[257:262] == b'ustar':
        return FileInfo('TAR', ('ZIP', 'TAR', 'GZIP'))
    if header.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        try:
            import olefile
            with olefile.OleFileIO(path) as ole:
                kind = 'XLS' if ole.exists('Workbook') or ole.exists('Book') else 'PPT' if ole.exists('PowerPoint Document') else 'DOC' if ole.exists('WordDocument') else None
            if kind:
                return FileInfo(kind, document_actions(kind), 'Benötigt LibreOffice.')
        except ImportError:
            return FileInfo('Office', (), 'Office-Erkennung benötigt olefile.')
    if header.startswith(b'{\\rtf'):
        return FileInfo('RTF', document_actions('RTF'), 'Benötigt LibreOffice.')
    # Plain text has no unique magic number. Validate UTF-8 and structure;
    # use the extension only to distinguish intentionally plain-text formats.
    if path.stat().st_size <= 20*1024*1024:
        try:
            text = path.read_text(encoding='utf-8-sig')
        except (UnicodeError, OSError):
            return None
        if '\x00' in text or not text.strip():
            return None
        if re.search(r'(?m)^(?:\d{1,2}:)?\d{2}:\d{2}[,.]\d{3}\s+-->\s+', text):
            kind = 'VTT' if text.lstrip().startswith('WEBVTT') else 'SRT'
            return FileInfo(kind, ('SRT', 'VTT', 'TXT'), 'Textbasierte Untertitel; keine Videoeinbettung.')
        ext = path.suffix.lower()
        kind = {'.md':'MD', '.markdown':'MD', '.txt':'TXT', '.html':'HTML', '.htm':'HTML', '.csv':'CSV', '.tsv':'TSV'}.get(ext)
        if kind in ('CSV', 'TSV'):
            delimiter = ',' if kind == 'CSV' else '\t'
            if kind == 'CSV':
                try:
                    csv.Sniffer().sniff(text[:8192], delimiters=',;')
                except csv.Error:
                    return FileInfo('TXT', document_actions('TXT'), 'Keine Tabellenstruktur erkannt.')
            elif delimiter not in text:
                return FileInfo('TXT', document_actions('TXT'), 'Keine Tabellenstruktur erkannt.')
        if kind:
            return FileInfo(kind, document_actions(kind), 'Benötigt LibreOffice oder Pandoc. Textformat anhand Inhalt und Endung erkannt.')
    return None

def convert_picture(source, target, output, svg=False, quality=92):
    checkpoint(0,"Bild wird verarbeitet …")
    if svg:
        from PySide6.QtSvg import QSvgRenderer
        from PySide6.QtGui import QImage, QPainter
        from PySide6.QtCore import Qt
        renderer = QSvgRenderer(svg_bytes(source))
        if not renderer.isValid():
            raise ValueError('SVG kann nicht gerendert werden.')
        size = renderer.defaultSize()
        if size.isEmpty() or size.width()*size.height() > 50_000_000:
            raise ValueError('SVG-Größe ungültig oder zu groß.')
        canvas = QImage(size, QImage.Format.Format_ARGB32)
        canvas.fill(Qt.GlobalColor.transparent)
        painter = QPainter(canvas); renderer.render(painter); painter.end()
        from PySide6.QtCore import QBuffer, QIODevice
        buffer = QBuffer(); buffer.open(QIODevice.OpenModeFlag.WriteOnly); canvas.save(buffer, 'PNG')
        original = Image.open(io.BytesIO(bytes(buffer.data())))
    else:
        original = Image.open(source)
    with original:
        if getattr(original, 'n_frames', 1) > 1:
            raise ValueError('Mehrseitige Bilder/Animationen noch nicht unterstützt.')
        image = ImageOps.exif_transpose(original)
        if target in ('JPG', 'BMP', 'PDF'):
            rgba = image.convert('RGBA'); image = Image.new('RGB', rgba.size, 'white'); image.paste(rgba, mask=rgba.getchannel('A'))
        elif image.mode not in ('RGB', 'RGBA'):
            image = image.convert('RGBA')
        fmt = {'JPG':'JPEG', 'HEIC':'HEIF'}.get(target, target)
        options = {'quality':quality} if target in ('JPG','WEBP','HEIC','AVIF') else {'resolution':300.0} if target == 'PDF' else {}
        image.save(output, format=fmt, **options)


def pdf_convert(source, target, output):
    import pymupdf
    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ValueError('Verschlüsseltes PDF wird nicht unterstützt.')
        if target in ('PNG', 'JPG'):
            with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
                for index, page in enumerate(document):
                    checkpoint(index/max(1,len(document))*100,'PDF-Seite wird exportiert …')
                    # Limit huge page raster allocations. Standard A4 is ~8.7 MP.
                    if page.rect.width*page.rect.height*(300/72)**2 > 80_000_000:
                        raise ValueError('PDF-Seite zu groß für 300-DPI-Export.')
                    image = page.get_pixmap(dpi=300, alpha=False)
                    archive.writestr(f'Seite_{index+1:04}.{target.lower()}', image.tobytes('jpeg' if target=='JPG' else 'png'))
            return
        texts=[]
        for index,page in enumerate(document):
            checkpoint(index/max(1,len(document))*100,'PDF-Text wird gelesen …')
            texts.append(page.get_text())
        if target == 'PPTX':
            from pptx import Presentation
            from pptx.util import Inches
            deck = Presentation()
            if len(document):
                first = document[0].rect
                deck.slide_width = Inches(first.width/72); deck.slide_height = Inches(first.height/72)
            for page in document:
                checkpoint(message='PDF-Folie wird erstellt …')
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                image = io.BytesIO(page.get_pixmap(dpi=150, alpha=False).tobytes('png'))
                slide.shapes.add_picture(image, 0, 0, width=deck.slide_width, height=deck.slide_height)
            deck.save(output)
            return
        if target == 'XLSX':
            from openpyxl import Workbook
            book = Workbook(); book.remove(book.active)
            count = 0
            for index, page in enumerate(document):
                checkpoint(index/max(1,len(document))*100,'Tabellen werden erkannt …')
                tables = page.find_tables()
                for table_index, table in enumerate(tables.tables):
                    sheet = book.create_sheet(f'Seite{index+1}_T{table_index+1}')
                    for row in table.extract():
                        sheet.append(row)
                        for cell in sheet[sheet.max_row]:
                            if cell.data_type == 'f':
                                cell.data_type = 's'
                    count += 1
            if not count:
                raise ValueError('Keine Tabellen erkannt. Keine OCR; PDF → XLSX benötigt erkennbare Tabellen.')
            book.save(output)
            return
        text = '\n\n'.join(texts)
        if not text.strip():
            raise ValueError('Kein Text gefunden; OCR ist noch nicht enthalten.')
        if target in ('TXT','MD'):
            Path(output).write_text(text, encoding='utf-8')
        elif target == 'DOCX':
            from docx import Document
            doc = Document()
            for index, text in enumerate(texts):
                if index: doc.add_page_break()
                for line in text.splitlines(): doc.add_paragraph(line)
            doc.save(output)
        elif target == 'EPUB':
            from ebooklib import epub
            from html import escape
            book = epub.EpubBook(); book.set_identifier(Path(source).stem); book.set_title(Path(source).stem); book.set_language('de')
            chapters = []
            for index, text in enumerate(texts):
                chapter = epub.EpubHtml(title=f'Seite {index+1}', file_name=f'page{index+1}.xhtml', lang='de')
                chapter.content = '<h1>Seite '+str(index+1)+'</h1><p>'+escape(text).replace('\n','<br/>')+'</p>'
                book.add_item(chapter); chapters.append(chapter)
            book.toc=tuple(chapters); book.spine=['nav']+chapters
            book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav()); epub.write_epub(str(output),book)
        else:
            raise ValueError('PDF-Zielformat nicht unterstützt.')


def office_convert(source, target, kind, output):
    pandoc = backend('pandoc')
    if kind in ('DOCX','MD','TXT','EPUB','HTML') and target in ('DOCX','MD','TXT','EPUB','HTML') and pandoc:
        reader = {'MD':'markdown', 'TXT':'markdown', 'DOCX':'docx', 'EPUB':'epub','HTML':'html'}[kind]
        writer = {'MD':'markdown', 'TXT':'plain', 'DOCX':'docx','EPUB':'epub','HTML':'html5'}[target]
        # Parse source in Pandoc's reader sandbox. Then remove resource-bearing
        # AST nodes before using system writer data (non-embedded Pandoc builds
        # cannot write DOCX/EPUB inside --sandbox).
        with tempfile.TemporaryDirectory() as temp:
            local = Path(temp)/('input'+source.suffix); shutil.copyfile(source, local)
            result = run([pandoc, '--sandbox', '-f', reader, '-t', 'json', str(local)], timeout=120)
            if result.returncode: raise ValueError(result.stderr[-1200:])
            ast = json.loads(result.stdout)
            def clean(value):
                if isinstance(value,list):
                    return [clean(item) for item in value]
                if isinstance(value,dict):
                    if value.get('t') == 'Image':
                        return {'t':'Str','c':'[Bild]'}
                    if value.get('t') == 'RawInline':
                        return {'t':'Str','c':''}
                    if value.get('t') == 'RawBlock':
                        return {'t':'Para','c':[]}
                    return {key:clean(item) for key,item in value.items()}
                return value
            ast = clean(ast); ast['meta'] = {}
            sanitized = Path(temp)/'safe.json'; sanitized.write_text(json.dumps(ast), encoding='utf-8')
            result = run([pandoc, '-f', 'json', '-t', writer, str(sanitized), '-o', str(output)], timeout=120)
            if result.returncode: raise ValueError(result.stderr[-1200:])
        return
    executable = soffice()
    if not executable:
        raise ValueError('LibreOffice fehlt.')
    filters = {'PDF':'pdf', 'DOCX':'docx:Office Open XML Text', 'TXT':'txt:Text (encoded):UTF8',
               'HTML':'html', 'EPUB':'epub', 'XLSX':'xlsx:Calc MS Excel 2007 XML', 'ODS':'ods',
               'CSV':'csv:Text - txt - csv (StarCalc):44,34,76,1',
               'TSV':'csv:Text - txt - csv (StarCalc):9,34,76,1'}
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp); outdir = root/'out'; outdir.mkdir(); profile = root/'profile'; profile.mkdir()
        # A private profile avoids an already running LibreOffice consuming the job.
        # Disable macro execution for this isolated profile.
        (profile/'user').mkdir()
        (profile/'user/registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>',encoding='utf-8')
        intermediate = 'PDF' if target in ('JPG','PNG') else target
        local = root/('input'+source.suffix); shutil.copyfile(source,local)
        result = run([executable, '-env:UserInstallation='+profile.as_uri(), '--headless','--norestore','--convert-to', filters[intermediate], '--outdir', str(outdir), str(local)],timeout=120)
        extension = 'csv' if intermediate == 'TSV' else intermediate.lower()
        candidate = outdir/('input.'+extension)
        if result.returncode or not candidate.is_file():
            raise ValueError('LibreOffice konnte diese Datei nicht umwandeln. '+result.stderr[-500:])
        if target in ('JPG','PNG'):
            pdf_convert(candidate,target,output)
        else:
            shutil.copyfile(candidate,output)


def subtitle_convert(source, target, output):
    text = source.read_text(encoding='utf-8-sig')
    lines = text.replace('\r\n','\n').splitlines()
    if target == 'TXT':
        kept = [line for line in lines if line.strip() and not line.strip().isdigit() and '-->' not in line and not line.startswith('WEBVTT')]
        result = '\n'.join(kept)
    elif target == 'VTT':
        body = re.sub(r'(\d{2}:\d{2}:\d{2}),(\d{3})',r'\1.\2',text)
        result = body if body.lstrip().startswith('WEBVTT') else 'WEBVTT\n\n'+body
    else:
        body = re.sub(r'^\ufeff?WEBVTT[^\n]*\n', '', text).strip()
        blocks = re.split(r'\n\s*\n',body); cues=[]
        for block in blocks:
            content=block.splitlines()
            timing=next((i for i,line in enumerate(content) if '-->' in line),None)
            if timing is None: continue
            # Drop WebVTT-only cue positioning metadata.
            match=re.match(r'((?:\d{2}:)?\d{2}:\d{2}[.,]\d{3})\s+-->\s+((?:\d{2}:)?\d{2}:\d{2}[.,]\d{3})',content[timing])
            if match:
                start, end = [value.replace('.',',') if value.count(':')==2 else '00:'+value.replace('.',',') for value in (match[1],match[2])]
                cues.append(str(len(cues)+1)+'\n'+start+' --> '+end+'\n'+'\n'.join(content[timing+1:]))
        if not cues: raise ValueError('Keine unterstützten Untertitel-Zeitmarken gefunden.')
        result='\n\n'.join(cues)+'\n'
    output.write_text(result,encoding='utf-8')


def archive_convert(source, target, kind, output):
    # Stream into a private normalized staging tree; never extract raw paths.
    # Reject links/devices/traversals and bound decompressed data.
    limit=2*1024**3; total=0; count=0
    with tempfile.TemporaryDirectory() as temp:
        root=Path(temp)
        seen=set()
        def store(name, stream, declared=0):
            nonlocal total,count
            name=name.replace('\\','/')
            parts=PurePosixPath(name).parts
            if not parts or any(p in ('..','') or ':' in p for p in parts) or name.startswith('/'):
                raise ValueError('Unsicherer Archivpfad: '+name)
            if any(p.rstrip(' .') != p or p.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))} for p in parts):
                raise ValueError('Unter Windows ungültiger Archivpfad: '+name)
            key='/'.join(parts).casefold()
            if key in seen: raise ValueError('Doppelter Archivpfad: '+name)
            seen.add(key); count+=1
            if count>10000 or total+declared>limit: raise ValueError('Archiv überschreitet das Limit (2 GiB / 10000 Dateien).')
            dest=root.joinpath(*parts); dest.parent.mkdir(parents=True,exist_ok=True)
            with dest.open('xb') as handle:
                while chunk:=stream.read(1024*1024):
                    checkpoint(message="Archiv wird verarbeitet …")
                    total+=len(chunk)
                    if total>limit: raise ValueError('Archiv überschreitet 2 GiB.')
                    handle.write(chunk)
        if kind=='ZIP':
            with zipfile.ZipFile(source) as archive:
                for member in archive.infolist():
                    if member.is_dir(): continue
                    mode=member.external_attr>>16
                    if mode & 0o170000 == 0o120000: raise ValueError('Archiv enthält symbolische Links.')
                    with archive.open(member) as handle: store(member.filename,handle,member.file_size)
        elif kind=='RAR':
            import rarfile
            with rarfile.RarFile(source) as archive:
                for member in archive.infolist():
                    if member.isdir(): continue
                    if member.is_symlink(): raise ValueError('Archiv enthält symbolische Links.')
                    with archive.open(member) as handle: store(member.filename,handle,member.file_size)
        elif kind in ('TAR','GZIP'):
            try:
                archive=tarfile.open(source,'r:*')
            except tarfile.ReadError:
                if kind!='GZIP': raise
                with gzip.open(source,'rb') as handle: store(source.stem,handle)
            else:
                with archive:
                    for member in archive:
                        if member.isdir(): continue
                        if not member.isfile(): raise ValueError('Nur reguläre Archivdateien unterstützt.')
                        with archive.extractfile(member) as handle: store(member.name,handle,member.size)
        if target=='ZIP':
            with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
                for file in root.rglob('*'):
                    checkpoint(message='ZIP wird geschrieben …')
                    if file.is_file():
                        with file.open('rb') as source_file, archive.open(str(file.relative_to(root)),'w',force_zip64=True) as target_file:
                            copy_stream(source_file,target_file)
        else:
            with tarfile.open(output,'w:gz' if target=='GZIP' else 'w') as archive:
                for file in root.rglob('*'):
                    checkpoint(message='TAR wird geschrieben …')
                    if file.is_file(): archive.add(file,arcname=str(file.relative_to(root)),recursive=False)
